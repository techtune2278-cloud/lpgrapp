"""Generate a PRIVATE, atomic Supabase import. Never commit its outputs.
Usage: python scripts/prepare-workbook.py workbook.xlsx output-directory
Requires openpyxl. Does not connect to any remote service.
"""
import sys, json, hashlib, re, unicodedata, collections
from pathlib import Path
from datetime import date, datetime
from decimal import Decimal
import openpyxl


def normal(value):
    return re.sub(r"\s+", " ", "".join(c for c in unicodedata.normalize("NFD", str(value or "").strip().lower().replace("’", "'")) if not unicodedata.combining(c)))


def number(value):
    if isinstance(value, bool) or value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return value
    return None  # textual guesses must never turn into financial entries


def encode(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def dumps(value):
    return json.dumps(value, ensure_ascii=False, default=encode, separators=(",", ":"), allow_nan=False)


def prepare(source):
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    formulas = openpyxl.load_workbook(source, data_only=False)
    values = openpyxl.load_workbook(source, data_only=True)
    rows, resources, sheets = [], [], []
    sections = {"02 PARAMETRES": "parameters", "08 TARIFS POSTAUX": "postal_rates", "09 PROJECTIONS": "projections", "12 EDITION ISBN": "editorial", "cyber scribe": "editorial", "financement 2025": "finance", "financement 2026": "finance", "prix de vente": "prices", "frais de port": "postal_history", "kipèkoi": "finance"}
    for sheet in formulas:
        count, templates = 0, 0
        for raw in sheet:
            if not any(c.value is not None for c in raw):
                continue
            r = raw[0].row
            # Preserve all source cells, including formulas, links, comments and cached errors.
            fv = [c.value if c.data_type == 'f' else None for c in raw]
            vv = [values[sheet.title].cell(r, c.column).value for c in raw]
            last = max((i for i, (v, f) in enumerate(zip(vv, fv)) if v not in (None, '') or f), default=-1) + 1
            vv, fv = vv[:last], fv[:last]
            meaningful = any(c.data_type != 'f' and c.value not in (None, '') for c in raw)
            meta = {c.coordinate: {**({'comment': c.comment.text} if c.comment else {}), **({'hyperlink': c.hyperlink.target} if c.hyperlink else {})} for c in raw if c.comment or c.hyperlink}
            payload = {'sheet_name': sheet.title, 'row_number': r, 'values': vv, 'formulas': fv, 'metadata': meta, 'kind': 'source' if meaningful else 'template'}
            payload['hash'] = hashlib.sha256(dumps(payload).encode()).hexdigest()
            rows.append(payload)
            count += 1
            templates += not meaningful
            if sheet.title in sections and meaningful:
                resources.append({'section': sections[sheet.title], 'sheet_name': sheet.title, 'row_number': r, 'values': vv})
        header_row = next((i for i in range(1,8) if normal(sheet.cell(i,1).value)=='id commande'),1) if sheet.title == '04 VENTES' else 4 if sheet.title[:2].isdigit() else 1
        labels = [str(c.value or '') if c.data_type != 'f' else '' for c in sheet[header_row]]
        sheets.append({'name': sheet.title, 'rows': count, 'template_rows': templates, 'headers': labels})

    sale_sheet = values['04 VENTES']
    header = next(r for r in range(1, 8) if normal(sale_sheet.cell(r, 1).value) == 'id commande')
    groups = collections.defaultdict(list)
    for r in range(header + 1, sale_sheet.max_row + 1):
        if sale_sheet.cell(r, 1).value not in (None, ''):
            groups[str(sale_sheet.cell(r, 1).value).strip()].append((r, [sale_sheet.cell(r, c).value for c in range(1, 19)]))
    orders, seen = [], collections.defaultdict(list)
    for ref, entries in groups.items():
        names = list(dict.fromkeys(str(v[2]).strip() for _, v in entries if v[2]))
        channels = list(dict.fromkeys(str(v[4]).strip() for _, v in entries if v[4]))
        amounts = [number(v[12]) for _, v in entries if number(v[12]) is not None]
        flags = []
        if len(names) != 1: flags.append('Plusieurs clients ou client absent pour cette référence')
        dates = list(dict.fromkeys(encode(v[1]) for _, v in entries if isinstance(v[1], (date, datetime))))
        if not dates: flags.append('Date de vente inconnue')
        if len(dates) > 1: flags.append('Plusieurs dates de vente')
        if not amounts: flags.append('Montant encaissé absent')
        elif len([a for a in amounts if a > 0]) > 1: flags.append('Plusieurs montants positifs : vérifier addition ou répétition')
        combined = normal(' '.join(str(v[i] or '') for _, v in entries for i in [4, 17]))
        classification = 'subscription' if 'abonnement' in combined else 'consignment' if 'depot' in combined else 'loan' if 'pret ' in combined else 'donation' if any(k in combined for k in ['offert', 'gratuit', 'don ']) else 'unclassified'
        if classification == 'unclassified' and amounts and sum(amounts) > 0: classification = 'sale'
        lines = []
        for r, v in entries:
            q, tome = number(v[7]), number(v[5])
            if q is None or q < 0 or int(q) != q: flags.append(f'Ligne {r} : quantité manquante ou invalide'); q = None
            if tome is None or int(tome) != tome or not 1 <= tome <= 33: flags.append(f'Ligne {r} : tome à préciser'); tome = None
            formula_error = any(c.data_type == 'f' and '#REF!' in c.value for c in formulas['04 VENTES'][r])
            if formula_error: flags.append(f'Ligne {r} : formule Excel cassée')
            if v[9] not in (None, '') and number(v[9]) is None: flags.append(f'Ligne {r} : prix unitaire non numérique')
            fingerprint = dumps([v[i] for i in [0, 1, 2, 3, 4, 5, 7, 8, 11, 12, 13, 14, 16, 17]])
            seen[fingerprint].append((ref, r))
            lines.append({'row_number': r, 'tome_number': tome, 'quantity': q, 'unit_price': number(v[9]), 'line_total': number(v[10]), 'amount_received': number(v[12]), 'payment_method': v[13], 'note': v[17], 'values': v})
        orders.append({'reference': ref, 'customer_name': ' / '.join(names) or 'Client non renseigné', 'channel': ' / '.join(channels) or None, 'sale_date': dates[0][:10] if len(dates) == 1 else None, 'source_received': round(sum(amounts), 2) if amounts else None, 'classification': classification, 'flags': list(dict.fromkeys(flags)), 'lines': lines})
    duplicates = [entries for entries in seen.values() if len(entries) > 1]
    for entries in duplicates:
        for ref, r in entries:
            order = next(o for o in orders if o['reference'] == ref)
            order['flags'].append(f'Ligne {r} : contenu métier répété, à rapprocher')

    stocks, catalogue = [], []
    for r in range(5, 38):
        s = values['03 CATALOGUE STOCKS']
        tome = number(s.cell(r, 1).value)
        if tome is None: continue
        initial, reprints = number(s.cell(r, 7).value), number(s.cell(r, 8).value)
        quantity = sum(l['quantity'] or 0 for o in orders for l in o['lines'] if l['tome_number'] == tome)
        # Compute independently; cached spreadsheet formulas may be stale.
        stocks.append({'tome_number': tome, 'initial_quantity': initial, 'reprint_quantity': reprints, 'source_quantity': quantity, 'theoretical_quantity': (initial or 0) + (reprints or 0) - quantity, 'physical_quantity': number(values['00 TABLEAU DE BORD'].cell(r, 9).value) if r < 16 else None})
        catalogue.append({'tome_number': tome, 'initial_quantity': initial})
    editorial = []
    for r in range(5, 16):
        v = [values['12 EDITION ISBN'].cell(r, c).value for c in range(1, 15)]
        if number(v[0]) is not None:
            editorial.append({'tome_number': v[0], 'subtitle': v[12], 'publisher': v[9], 'source': v})
    expenses = []
    for r in range(5, values['06 DEPENSES'].max_row + 1):
        v = [values['06 DEPENSES'].cell(r, c).value for c in range(1, 12)]
        if v[0]: expenses.append({'reference': str(v[0]), 'amount': number(v[5]), 'date': encode(v[1]), 'category': v[2], 'object': v[3], 'supplier': v[4], 'paid_by': v[6], 'status': v[9], 'note': v[10]})
    summary = {'sheet_count': len(sheets), 'source_rows': len(rows), 'order_count': len(orders), 'sale_line_count': sum(len(o['lines']) for o in orders), 'source_quantity': sum(l['quantity'] or 0 for o in orders for l in o['lines']), 'source_received': round(sum(o['source_received'] or 0 for o in orders), 2), 'duplicate_groups': len(duplicates), 'physical_counts': sum(s['physical_quantity'] is not None for s in stocks), 'physical_total': sum(s['physical_quantity'] or 0 for s in stocks), 'expense_total': round(sum(e['amount'] or 0 for e in expenses), 2)}
    return {'format_version': 1, 'source_name': source.name, 'sha256': digest, 'sheets': sheets, 'summary': summary, 'rows': rows, 'orders': orders, 'stocks': stocks, 'catalogue': catalogue, 'editorial': editorial, 'expenses': expenses, 'resources': resources}


def main():
    source, dest = Path(sys.argv[1]), Path(sys.argv[2])
    dest.mkdir(parents=True, exist_ok=True)
    payload = prepare(source)
    text = dumps(payload)
    (dest / 'reprise-excel-privee.json').write_text(text, encoding='utf-8')
    # Dollar quote is chosen from the generated contents, not interpolated user SQL.
    delimiter = '$workbook_' + payload['sha256'][:20] + '$'
    assert delimiter not in text
    (dest / '02_IMPORT_PRIVE_SUPABASE.sql').write_text('-- CONFIDENTIEL : ne pas publier sur GitHub ou Vercel.\n-- Exécuter APRÈS les migrations 006 et 007. Aucun stock actif ne sera changé.\nBEGIN;\nSELECT public.import_workbook_snapshot(' + delimiter + text + delimiter + '::jsonb);\nCOMMIT;\n', encoding='utf-8')
    (dest / 'CONTROLES_SANS_COORDONNEES.json').write_text(json.dumps(payload['summary'], indent=2, ensure_ascii=False), encoding='utf-8')
    print(json.dumps(payload['summary'], ensure_ascii=False))


if __name__ == '__main__':
    main()
